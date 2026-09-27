"""
Candidate Signal Prep — Claude Code rebuild
===========================================
Rebuild of the n8n Code node "Candidate Signal Prep" from the workflow
"Retail Anomaly Agentic Workflow - v3 (Risk Classifier + Alert Drafting)".

Input : list of transaction dicts + list of member dicts
        (same columns as the Google Sheets sources in n8n).
Output: one dict per transaction with 6 rule flags (R1-R6), evidence text,
        anomaly_count, is_anomaly, verified_at — same field names as n8n,
        so the output can be written back to the same verification sheet.

Rules
  R1 card sharing       : same card (last4 + issuing bank) used by >1 member
  R2 tier mismatch      : item tier requirement above member tier
  R3 outside hours      : transaction hour < 9 or > 21
  R4 bulk limited item  : Platinum item qty > 1, Gold item qty > 2
  R5 impossible travel  : same member (GUEST excluded), two physical stores in
                          different geo_area within 90 minutes
  R6 duplicate account  : member shares phone or email with another member
"""

from __future__ import annotations

import csv
import re
import sys
from collections import defaultdict
from datetime import datetime, timezone

TIER_RANK = {"regular": 0, "silver": 1, "gold": 2, "platinum": 3}
OPERATING_HOUR_START = 9
OPERATING_HOUR_END = 21
IMPOSSIBLE_TRAVEL_MINUTES = 90
QTY_LIMIT = {"platinum": 1, "gold": 2}
GUEST_ID = "GUEST"


# ---------- helpers ----------------------------------------------------------
def _rank(tier) -> int:
    if tier is None:
        return -1
    return TIER_RANK.get(str(tier).strip().lower(), -1)


def _is_true(v) -> bool:
    return v is True or str(v).strip().lower() == "true"


def _key(v) -> str:
    return "" if v is None else str(v).strip()


def card_key_of(t: dict) -> str | None:
    """Card identity = last 4 digits + issuing bank (payment_type excluded)."""
    if not _is_true(t.get("is_card_payment")):
        return None
    last4 = re.sub(r"\D", "", str(t.get("card_last4") or ""))
    bank = str(t.get("issuing_bank") or "").strip().upper()
    if not last4 or not bank:
        return None
    return f"{last4}|{bank}"


def parse_dt(value) -> datetime | None:
    """Accepts the sheet format '4/28/26 16:12' plus ISO strings."""
    if not value:
        return None
    s = str(value).strip()
    for fmt in ("%m/%d/%y %H:%M", "%m/%d/%Y %H:%M", "%Y-%m-%d %H:%M:%S",
                "%Y-%m-%d %H:%M", "%Y-%m-%dT%H:%M:%S"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    return None


def _is_online(t: dict, city) -> bool:
    return (str(t.get("channel") or "").strip().lower() == "online"
            or str(city or "").strip().lower() == "online")


# ---------- cross-transaction indexes ----------------------------------------
def build_duplicate_evidence(members: list[dict]) -> dict[str, list[str]]:
    phone, email = defaultdict(list), defaultdict(list)
    for m in members:
        mid = _key(m.get("member_id"))
        if m.get("phone_number"):
            phone[str(m["phone_number"]).strip()].append(mid)
        if m.get("email"):
            email[str(m["email"]).strip().lower()].append(mid)

    evidence: dict[str, list[str]] = defaultdict(list)
    for label, groups in (("nomor telepon", phone), ("email", email)):
        for ids in groups.values():
            if len(ids) > 1:
                for mid in ids:
                    others = [x for x in ids if x != mid]
                    evidence[mid].append(
                        f"{label} sama dengan member {', '.join(others)}")
    return evidence


def build_card_groups(transactions: list[dict]) -> dict[str, list[str]]:
    groups: dict[str, list[str]] = defaultdict(list)
    for t in transactions:
        ck = card_key_of(t)
        mid = _key(t.get("member_id"))
        if ck and mid and mid not in groups[ck]:
            groups[ck].append(mid)
    return groups


def _area(t: dict) -> str:
    """Area pembanding R5: geo_area (Jakarta Selatan/Utara/Pusat = Jakarta), fallback store_city."""
    return str(t.get("geo_area") or t.get("store_city") or "").strip()


def build_travel_evidence(transactions: list[dict]) -> dict[str, str]:
    by_member: dict[str, list[dict]] = defaultdict(list)
    for t in transactions:
        mid = _key(t.get("member_id"))
        # GUEST bukan satu orang: semua non-member berbagi ID yang sama.
        if mid and mid.upper() != GUEST_ID:
            by_member[mid].append(t)

    evidence: dict[str, str] = {}
    for rows in by_member.values():
        rows = sorted(rows, key=lambda r: parse_dt(r.get("transaction_datetime")) or datetime.min)
        for prev, curr in zip(rows, rows[1:]):
            pc = prev.get("store_city") or prev.get("geo_area")
            cc = curr.get("store_city") or curr.get("geo_area")
            if _is_online(prev, pc) or _is_online(curr, cc):
                continue
            pa, ca = _area(prev), _area(curr)
            if not pa or not ca or pa.lower() == ca.lower():
                continue
            d0, d1 = parse_dt(prev.get("transaction_datetime")), parse_dt(curr.get("transaction_datetime"))
            if not d0 or not d1:
                continue
            diff = (d1 - d0).total_seconds() / 60
            if 0 <= diff < IMPOSSIBLE_TRAVEL_MINUTES:
                text = (f"{pc} pada {prev['transaction_datetime']} lalu {cc} pada "
                        f"{curr['transaction_datetime']} (selisih {round(diff)} menit)")
                evidence[curr["transaction_id"]] = text
                evidence[prev["transaction_id"]] = text
    return evidence


# ---------- main step --------------------------------------------------------
def candidate_signal_prep(transactions: list[dict], members: list[dict]) -> list[dict]:
    member_by_id = {_key(m["member_id"]): m for m in members if m.get("member_id")}
    dup_ev = build_duplicate_evidence(members)
    card_groups = build_card_groups(transactions)
    travel_ev = build_travel_evidence(transactions)
    verified_at = datetime.now(timezone.utc).isoformat(timespec="milliseconds").replace("+00:00", "Z")

    out = []
    for t in transactions:
        mid = _key(t.get("member_id"))

        # R1 card sharing
        ck = card_key_of(t)
        sharers = [x for x in card_groups.get(ck, []) if x != mid] if ck else []
        r1 = bool(sharers)
        if not ck:
            ev1 = "Bukan transaksi kartu atau data kartu (last4/bank) tidak lengkap - R1 tidak dinilai"
        elif r1:
            ev1 = (f"Kartu {t.get('issuing_bank') or ''} {t.get('card_last4') or ''} juga dipakai oleh member "
                   f"{', '.join(sharers)} (key: last4 + bank, perlu verifikasi karena last4 tidak unik)")
        else:
            ev1 = "Tidak ada indikasi kartu dipakai member lain pada batch ini"

        # R2 tier mismatch (member master tier takes precedence)
        m_tier = (member_by_id.get(mid) or {}).get("member_tier") or t.get("member_tier")
        r2 = _rank(t.get("item_tier_requirement")) > _rank(m_tier)
        ev2 = (f"Item membutuhkan tier {t.get('item_tier_requirement')} namun member bertier {m_tier}"
               if r2 else "Tier item sesuai atau di bawah tier member")

        # R3 outside operating hours
        hour = t.get("transaction_hour")
        if hour in (None, ""):
            d = parse_dt(t.get("transaction_datetime"))
            hour = d.hour if d else None
        hour = int(hour) if hour not in (None, "") else None
        r3 = hour is not None and (hour < OPERATING_HOUR_START or hour > OPERATING_HOUR_END)
        ev3 = (f"Transaksi pukul {hour}:00, di luar asumsi jam operasional "
               f"{OPERATING_HOUR_START}-{OPERATING_HOUR_END}"
               if r3 else "Transaksi dalam rentang jam operasional yang diasumsikan")

        # R4 bulk purchase of limited items
        tier_key = str(t.get("item_tier_requirement") or "").strip().lower()
        try:
            qty = float(t.get("qty") or 0)
        except ValueError:
            qty = 0
        qty = int(qty) if qty == int(qty) else qty
        limit = QTY_LIMIT.get(tier_key)
        r4 = limit is not None and qty > limit
        ev4 = (f"Item {tier_key.capitalize()} dibeli {qty} pcs, melebihi batas {limit}"
               if r4 else "Kuantitas dalam batas wajar")

        # R5 impossible travel
        r5 = t.get("transaction_id") in travel_ev
        ev5 = travel_ev.get(t.get("transaction_id"),
                            "Tidak ada transaksi lain dari member ini di kota berbeda dalam waktu singkat")

        # R6 duplicate account
        r6 = bool(mid and dup_ev.get(mid))
        ev6 = "; ".join(dup_ev[mid]) if r6 else "Tidak ada kontak (telepon/email) yang sama dengan member lain"

        count = sum([r1, r2, r3, r4, r5, r6])
        row = dict(t)
        row.update({
            "card_identity_key": ck or "",
            "flag_R1_card_sharing": r1, "evidence_R1": ev1,
            "flag_R2_tier_mismatch": r2, "evidence_R2": ev2,
            "flag_R3_outside_hours": r3, "evidence_R3": ev3,
            "flag_R4_bulk_limited": r4, "evidence_R4": ev4,
            "flag_R5_impossible_travel": r5, "evidence_R5": ev5,
            "flag_R6_duplicate_account": r6, "evidence_R6": ev6,
            "anomaly_count": count,
            "is_anomaly": count > 0,
            "is_anomaly_num": 1 if count > 0 else 0,
            "verified_at": verified_at,
        })
        out.append(row)
    return out


# ---------- CLI: python candidate_signal_prep.py trx.csv members.csv|xlsx out.csv
def _read_csv(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        # Ekspor Google Sheets membawa kolom kosong tanpa header; buang agar tidak ikut ke output.
        return [{k: v for k, v in row.items() if k} for row in csv.DictReader(f)]


def _read_xlsx(path):
    """Sheet pertama, semua nilai dijadikan teks agar sama dengan hasil baca CSV."""
    from openpyxl import load_workbook  # opsional: hanya dibutuhkan untuk input .xlsx

    ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    return [{h: "" if v is None else str(v) for h, v in zip(header, r) if h}
            for r in rows if any(v is not None for v in r)]


def read_table(path):
    return _read_xlsx(path) if str(path).lower().endswith((".xlsx", ".xlsm")) else _read_csv(path)


if __name__ == "__main__":
    if len(sys.argv) != 4:
        sys.exit("Usage: python candidate_signal_prep.py transactions.csv members.csv|members.xlsx output.csv")
    rows = candidate_signal_prep(read_table(sys.argv[1]), read_table(sys.argv[2]))
    with open(sys.argv[3], "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)
    flagged = sum(r["is_anomaly_num"] for r in rows)
    print(f"{len(rows)} transaksi diproses, {flagged} terindikasi anomali -> {sys.argv[3]}")
